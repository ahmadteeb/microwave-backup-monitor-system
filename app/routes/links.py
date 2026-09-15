import os
import csv
import io
import re
import uuid
import logging
import mimetypes
import openpyxl
from datetime import datetime
from werkzeug.utils import secure_filename
from flask import Blueprint, request, jsonify, Response, session, send_file, current_app, send_from_directory
from app.models import db, Link, PingResult, AppSettings, LinkStatus, LinkAttachment
from app.services.ping_service import ping_single_link
from app.services.notification_service import send_event_notification
from app.services.external_util_service import refresh_external_utilization_for_single_link, lookup_link_info, lookup_leg_info
from app.services.log_service import write_log
from app.permissions import login_required, require_permission

logger = logging.getLogger(__name__)

links_bp = Blueprint('links', __name__, url_prefix='/api/links')

def validate_ipv4(ip):
    pattern = re.compile(r'^((25[0-5]|(2[0-4]|1\d|[1-9]|)\d)\.?\b){4}$')
    return pattern.match(ip)

def get_file_type(filename):
    ext = filename.rsplit('.', 1)[-1].lower() if '.' in filename else ''
    if ext == 'pdf':
        return 'pdf'
    if ext in ('vsd', 'vsdx', 'vssx', 'vstx'):
        return 'visio'
    return 'other'

def is_allowed_attachment(filename):
    if '.' not in filename:
        return False
    ext = filename.rsplit('.', 1)[1].lower()
    allowed = current_app.config.get('ALLOWED_ATTACHMENT_EXTENSIONS', {'pdf', 'vsd', 'vsdx', 'vssx', 'vstx'})
    return ext in allowed

def save_attachment_file(file_storage, link_id, user_id=None):
    if not file_storage or not file_storage.filename:
        return None
    
    orig_name = file_storage.filename
    if not is_allowed_attachment(orig_name):
        raise ValueError(f"File type not allowed for '{orig_name}'. Allowed: PDF (.pdf), Visio (.vsd, .vsdx)")
    
    clean_name = secure_filename(orig_name)
    if not clean_name:
        clean_name = "attachment"
    
    # Enforce only one file per link: remove any existing attachments
    existing_atts = LinkAttachment.query.filter_by(link_id=link_id).all()
    for old_att in existing_atts:
        remove_attachment_file_from_disk(old_att.filename)
        db.session.delete(old_att)
    
    stored_name = f"{uuid.uuid4().hex}_{clean_name}"
    upload_dir = current_app.config.get('UPLOAD_FOLDER', os.path.join(current_app.config.get('APP_ROOT', '.'), 'data', 'uploads', 'attachments'))
    os.makedirs(upload_dir, exist_ok=True)
    
    file_path = os.path.join(upload_dir, stored_name)
    file_storage.save(file_path)
    file_size = os.path.getsize(file_path)
    
    attachment = LinkAttachment(
        link_id=link_id,
        filename=stored_name,
        original_filename=orig_name,
        file_type=get_file_type(orig_name),
        file_size=file_size,
        uploaded_by_id=user_id
    )
    db.session.add(attachment)
    return attachment


def remove_attachment_file_from_disk(stored_filename):
    try:
        upload_dir = current_app.config.get('UPLOAD_FOLDER', os.path.join(current_app.config.get('APP_ROOT', '.'), 'data', 'uploads', 'attachments'))
        if upload_dir and stored_filename:
            path = os.path.join(upload_dir, stored_filename)
            if os.path.exists(path):
                import gc
                import time
                for _ in range(3):
                    try:
                        os.remove(path)
                        break
                    except (PermissionError, OSError):
                        gc.collect()
                        time.sleep(0.05)
    except Exception as e:
        logger.warning(f"Failed to remove attachment file {stored_filename} from disk: {e}")


def serialize_link(link):
    # Get latest ping result
    latest_ping = link.ping_results.order_by(PingResult.timestamp.desc()).first()
    
    status = 'UNKNOWN'
    if link.status:
        status = link.status.mw_status.upper()

    ping_data = None
    latency = None
    if latest_ping:
        ping_data = {
            "reachable": latest_ping.reachable,
            "latency_ms": latest_ping.latency_ms,
            "packet_loss": latest_ping.packet_loss,
            "timestamp": latest_ping.timestamp.isoformat() + "Z"
        }
        latency = latest_ping.latency_ms
        if status == 'UNKNOWN':
            if latest_ping.reachable:
                status = 'UP'
            else:
                if latest_ping.raw_output and "100% packet loss" in latest_ping.raw_output:
                    status = 'TIMEOUT'
                else:
                    status = 'DOWN'

    metric_data = None
    leg_util_pct = None
    leg_bitrate = None
    leg_capacity_pct = None
    
    if link.status:
        metric_data = {
            "leg_util_pct": link.status.leg_util_pct,
            "leg_capacity_mbps": link.status.leg_capacity_mbps,
            "mw_util_pct": link.status.mw_util_pct,
            "mw_capacity_mbps": link.status.mw_capacity_mbps,
            "timestamp": link.status.last_metric_at.isoformat() + "Z" if link.status.last_metric_at else None
        }
        settings = db.session.get(AppSettings, 1)
        warn_pct = settings.util_warning_threshold_pct if settings else 70.0
        if status == 'UP' and link.status.mw_util_pct is not None and link.status.mw_util_pct >= warn_pct:
            status = 'HIGH'

        leg_util_pct = link.status.leg_util_pct
        if link.status.avg_max_mbitrate:
            leg_bitrate = link.status.avg_max_mbitrate
            if link.status.mw_capacity_mbps:
                try:
                    leg_capacity_pct = round((float(link.status.mw_capacity_mbps) / float(link.status.avg_max_mbitrate)) * 100, 1)
                except (TypeError, ValueError, ZeroDivisionError):
                    pass

    attachments_list = []
    if hasattr(link, 'attachments'):
        for att in link.attachments.all():
            attachments_list.append({
                "id": att.id,
                "filename": att.original_filename,
                "file_type": att.file_type,
                "file_size": att.file_size,
                "uploaded_at": att.uploaded_at.isoformat() + "Z" if att.uploaded_at else None,
                "is_pdf": att.file_type == 'pdf' or att.original_filename.lower().endswith('.pdf'),
                "is_visio": att.file_type == 'visio' or att.original_filename.lower().endswith(('.vsd', '.vsdx', '.vssx', '.vstx'))
            })

    return {
        "id": link.id,
        "link_id": link.link_id,
        "leg_name": link.leg_name,
        "site_a": link.site_a,
        "site_b": link.site_b,
        "mw_ip": link.mw_ip,
        "link_type": link.link_type,
        "notes": link.notes,
        "status": status,
        "leg_util_pct": leg_util_pct,
        "leg_bitrate": leg_bitrate,
        "leg_capacity_pct": leg_capacity_pct,
        "latency_ms": latency,
        "latest_ping": ping_data,
        "latest_metric": metric_data,
        # Per-link utilization thresholds (Tier 3.2)
        "util_warning_threshold_pct": link.util_warning_threshold_pct,
        "util_critical_threshold_pct": link.util_critical_threshold_pct,
        # Attachments
        "attachments": attachments_list
    }


@links_bp.route('', methods=['GET'])
@login_required
@require_permission('links.view')
def list_links():
    status_filter = request.args.get('status')
    leg_filter = request.args.get('leg')
    search_query = request.args.get('search')
    page = request.args.get('page', 1, type=int)
    per_page = request.args.get('per_page', 20, type=int)

    query = Link.query

    if leg_filter and leg_filter != 'ALL_REGIONS':
        query = query.filter(Link.leg_name == leg_filter)
    
    if search_query:
        search_term = f"%{search_query}%"
        query = query.filter(db.or_(
            Link.link_id.ilike(search_term),
            Link.leg_name.ilike(search_term)
        ))

    if status_filter and status_filter != 'ALL_OPERATIONAL':
        status_lower = status_filter.lower()
        query = query.outerjoin(LinkStatus, Link.id == LinkStatus.link_id)
        query = query.filter(LinkStatus.mw_status == status_lower)

    paginated = query.paginate(page=page, per_page=per_page, error_out=False)
    results = [serialize_link(link) for link in paginated.items]

    return jsonify({
        "links": results,
        "total": paginated.total,
        "page": paginated.page,
        "per_page": paginated.per_page,
        "pages": paginated.pages
    })

@links_bp.route('/export', methods=['GET'])
@login_required
@require_permission('links.export')
def export_links():
    status_filter = request.args.get('status')
    leg_filter = request.args.get('leg')
    search_query = request.args.get('search')

    query = Link.query
    if leg_filter and leg_filter != 'ALL_REGIONS':
        query = query.filter(Link.leg_name == leg_filter)
    if search_query:
        search_term = f"%{search_query}%"
        query = query.filter(db.or_(
            Link.link_id.ilike(search_term),
            Link.leg_name.ilike(search_term)
        ))
    if status_filter and status_filter != 'ALL_OPERATIONAL':
        status_lower = status_filter.lower()
        query = query.outerjoin(LinkStatus, Link.id == LinkStatus.link_id)
        query = query.filter(LinkStatus.mw_status == status_lower)

    links = query.all()
    exported = [serialize_link(link) for link in links]

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Link ID", "Diagram Attachment", "Leg", "Site A", "Site B", "LEG Util %", "LEG Bitrate Mbps", "MW IP", "Link Type", "Status", "Latency ms", "Direct LEG Util %", "MW Util %", "MW/LEG %", "Link Cap Mbps", "Last Ping"])

    for link in exported:
        diagram_name = link["attachments"][0]["filename"] if link.get("attachments") and len(link["attachments"]) > 0 else ''
        writer.writerow([
            link["link_id"],
            diagram_name,
            link["leg_name"],
            link["site_a"],
            link["site_b"],
            link["leg_util_pct"] if link["leg_util_pct"] is not None else '',
            link["leg_bitrate"] if link["leg_bitrate"] is not None else '',
            link["mw_ip"],
            link["link_type"],
            link["status"],
            link["latency_ms"] if link["latency_ms"] is not None else '',
            link["latest_metric"]["leg_util_pct"] if link["latest_metric"] else '',
            link["latest_metric"]["mw_util_pct"] if link["latest_metric"] else '',
            link["leg_capacity_pct"] if link["leg_capacity_pct"] is not None else '',
            link["latest_metric"]["mw_capacity_mbps"] if link["latest_metric"] else '',
            link["latest_ping"]["timestamp"] if link["latest_ping"] else ''
        ])

    return Response(output.getvalue(), mimetype='text/csv', headers={
        'Content-Disposition': 'attachment; filename="active_link_inventory.csv"'
    })

def _extract_link_payload():
    if request.is_json:
        return request.get_json() or {}
    
    data = {}
    for key, value in request.form.items():
        if key in ('util_warning_threshold_pct', 'util_critical_threshold_pct'):
            if value is not None and str(value).strip() != '':
                try:
                    data[key] = float(value)
                except ValueError:
                    data[key] = None
            else:
                data[key] = None
        else:
            data[key] = value
    return data

def _process_uploaded_files(link_id, user_id=None):
    f = None
    if 'attachment' in request.files and request.files['attachment'].filename:
        f = request.files['attachment']
    elif 'file' in request.files and request.files['file'].filename:
        f = request.files['file']
    else:
        files = request.files.getlist('attachments') or request.files.getlist('files')
        if files and len(files) > 0 and files[0].filename:
            f = files[0]
            
    if f and f.filename and f.filename.strip():
        att = save_attachment_file(f, link_id, user_id=user_id)
        return [att] if att else []
    return []


@links_bp.route('', methods=['POST'])
@login_required
@require_permission('links.add')
def create_link():
    data = _extract_link_payload()
    if not data or not data.get('link_id') or not data.get('leg_name') or not data.get('mw_ip'):
        return jsonify({"error": "Missing required fields: link_id, leg_name, mw_ip"}), 400
    
    if not validate_ipv4(data['mw_ip']):
        return jsonify({"error": "Invalid IPv4 address format"}), 400

    existing = Link.query.filter_by(link_id=data['link_id']).first()
    if existing:
        return jsonify({"error": f"Link ID {data['link_id']} already exists"}), 409

    user_id = session.get('user_id')
    link = Link(
        link_id=data['link_id'],
        leg_name=data['leg_name'],
        site_a=data.get('site_a'),
        site_b=data.get('site_b'),
        mw_ip=data['mw_ip'],
        link_type=data.get('link_type', 'microwave'),
        notes=data.get('notes'),
        is_active=True,
        # Per-link thresholds (Tier 3.2)
        util_warning_threshold_pct=data.get('util_warning_threshold_pct'),
        util_critical_threshold_pct=data.get('util_critical_threshold_pct'),
        created_by_id=user_id
    )
    db.session.add(link)
    db.session.flush()

    # Process any attached Visio or PDF files (required)
    try:
        saved_files = _process_uploaded_files(link.id, user_id=user_id)
    except ValueError as val_err:
        db.session.rollback()
        return jsonify({"error": str(val_err)}), 400

    if not saved_files or len(saved_files) == 0:
        db.session.rollback()
        return jsonify({"error": "Link diagram attachment is required. Please attach a Visio (.vsd, .vsdx) or PDF file."}), 400

    db.session.commit()


    # Audit log: link created (Tier 2.4)
    write_log('links', 'link_created', session.get('username', 'system'), link.link_id,
              {'leg_name': link.leg_name, 'mw_ip': link.mw_ip, 'attachments': len(saved_files)}, ip_address=request.remote_addr)
    
    # Fetch initial data from external DB
    try:
        refresh_external_utilization_for_single_link(link)
        db.session.commit()
    except Exception as e:
        logger.warning(f"Failed to fetch initial external status for {link.link_id}: {e}")

    return jsonify(serialize_link(link)), 201

def _resolve_external_capacity(row):
    capacity = row.get('XPIC_MW_Link_Capacity')
    if capacity is None:
        capacity = row.get('MW_Link_Capacity')
    return capacity


def _persist_external_link_metrics(link, row):
    if row.get('Source_NE_Card'):
        link.site_a = row.get('Source_NE_Card')
    if row.get('Sink_NE_Card'):
        link.site_b = row.get('Sink_NE_Card')

    status = LinkStatus.query.filter_by(link_id=link.id).first()
    if not status:
        status = LinkStatus(link_id=link.id)
        db.session.add(status)
        
    status.mw_util_pct = row.get('AVG_MAX_Util_RxTx_perc')
    status.mw_capacity_mbps = _resolve_external_capacity(row)
    status.metric_source = 'external'
    status.last_metric_at = datetime.utcnow()

    link.updated_at = datetime.utcnow()
    db.session.commit()

@links_bp.route('/lookup', methods=['POST'])
@login_required
@require_permission('links.view')
def lookup_external_link():
    data = request.get_json() or {}
    link_id = data.get('link_id', '').strip()
    if not link_id:
        return jsonify({'error': 'Missing required field: link_id'}), 400

    try:
        row = lookup_link_info(link_id)
    except RuntimeError as exc:
        return jsonify({'error': str(exc)}), 503
    except Exception as exc:
        return jsonify({'error': f'External lookup failed: {exc}'}), 500

    if not row:
        return jsonify({'error': 'Link not found in external utilization database'}), 404

    link = Link.query.filter_by(link_id=link_id).first()
    if link:
        try:
            _persist_external_link_metrics(link, row)
        except Exception:
            db.session.rollback()

    return jsonify({'external': row}), 200

@links_bp.route('/lookup-leg', methods=['POST'])
@login_required
@require_permission('links.view')
def lookup_external_leg():
    data = request.get_json() or {}
    leg_name = data.get('leg_name', '').strip()
    if not leg_name:
        return jsonify({'error': 'Missing required field: leg_name'}), 400

    try:
        row = lookup_leg_info(leg_name)
    except RuntimeError as exc:
        return jsonify({'error': str(exc)}), 503
    except Exception as exc:
        return jsonify({'error': f'External lookup failed: {exc}'}), 500

    if not row:
        return jsonify({'error': 'LEG not found in external utilization database'}), 404

    try:
        links = Link.query.filter_by(leg_name=row.get('LEG_Name')).all()
        for link in links:
            status = LinkStatus.query.filter_by(link_id=link.id).first()
            if not status:
                status = LinkStatus(link_id=link.id)
                db.session.add(status)
            status.avg_max_mbitrate = row.get('AVG_MAX_MBitRate')
            status.interface_speed_min = row.get('Interface_Speed_Min')
            status.interface_speed_max = row.get('Interface_Speed_Max')
            status.sub_leg_count = row.get('Sub_LEG_Count')
            status.leg_source = 'external'
            
            if status.interface_speed_max:
                try:
                    status.leg_capacity_mbps = float(status.interface_speed_max)
                    status.leg_util_pct = round((float(status.avg_max_mbitrate or 0) / float(status.interface_speed_max)) * 100, 1)
                except (TypeError, ValueError, ZeroDivisionError):
                    pass
            
        db.session.commit()
    except Exception:
        db.session.rollback()

    return jsonify({'external': row}), 200

@links_bp.route('/<int:id>', methods=['GET'])
@login_required
@require_permission('links.view')
def get_link(id):
    link = db.session.get(Link, id)
    if not link:
        return jsonify({'error': 'Link not found'}), 404
    serialized = serialize_link(link)
    
    ping_history = PingResult.query.filter_by(link_id=id).order_by(PingResult.timestamp.desc()).limit(1440).all()
    
    return jsonify({
        "link": serialized,
        "ping_history": [{
            "timestamp": p.timestamp.isoformat() + "Z",
            "reachable": p.reachable,
            "latency_ms": p.latency_ms,
            "packet_loss": p.packet_loss
        } for p in ping_history],
        "metric_history": []
    }), 200

@links_bp.route('/<int:id>', methods=['PUT'])
@login_required
@require_permission('links.edit')
def update_link(id):
    link = db.session.get(Link, id)
    if not link:
        return jsonify({'error': 'Link not found'}), 404
    data = _extract_link_payload()
    
    if 'link_id' in data and data['link_id']:
        if data['link_id'] != link.link_id and Link.query.filter_by(link_id=data['link_id']).first():
            return jsonify({"error": f"Link ID {data['link_id']} already exists"}), 409
        link.link_id = data['link_id']

    if 'mw_ip' in data and data['mw_ip']:
        if not validate_ipv4(data['mw_ip']):
            return jsonify({"error": "Invalid IPv4 address format"}), 400
        link.mw_ip = data['mw_ip']

    if 'leg_name' in data: link.leg_name = data['leg_name']
    if 'site_a' in data: link.site_a = data['site_a']
    if 'site_b' in data: link.site_b = data['site_b']
    if 'link_type' in data: link.link_type = data['link_type']
    if 'notes' in data: link.notes = data['notes']

    # Per-link thresholds (Tier 3.2)
    if 'util_warning_threshold_pct' in data:
        link.util_warning_threshold_pct = data['util_warning_threshold_pct']
    if 'util_critical_threshold_pct' in data:
        link.util_critical_threshold_pct = data['util_critical_threshold_pct']

    user_id = session.get('user_id')
    # Process any attached Visio or PDF files
    try:
        saved_files = _process_uploaded_files(link.id, user_id=user_id)
    except ValueError as val_err:
        db.session.rollback()
        return jsonify({"error": str(val_err)}), 400

    # Ensure link has at least 1 diagram attachment
    if link.attachments.count() == 0 and (not saved_files or len(saved_files) == 0):
        db.session.rollback()
        return jsonify({"error": "Link diagram attachment is required. Please attach a Visio (.vsd, .vsdx) or PDF file."}), 400

    db.session.commit()


    # Refresh data from external DB in case link_id or leg_name changed
    try:
        refresh_external_utilization_for_single_link(link)
        db.session.commit()
    except Exception as e:
        logger.warning(f"Failed to refresh external status for {link.link_id}: {e}")

    return jsonify(serialize_link(link)), 200

@links_bp.route('/<int:id>/attachments', methods=['POST'])
@login_required
@require_permission('links.edit')
def upload_attachment(id):
    link = db.session.get(Link, id)
    if not link:
        return jsonify({'error': 'Link not found'}), 404
        
    user_id = session.get('user_id')
    try:
        saved_files = _process_uploaded_files(link.id, user_id=user_id)
    except ValueError as val_err:
        return jsonify({"error": str(val_err)}), 400

    if not saved_files:
        return jsonify({'error': 'No file uploaded or selected'}), 400

    db.session.commit()
    write_log('links', 'attachment_uploaded', session.get('username', 'system'), link.link_id,
              {'count': len(saved_files), 'filenames': [a.original_filename for a in saved_files]}, ip_address=request.remote_addr)

    return jsonify({
        'message': f'Uploaded {len(saved_files)} attachment(s)',
        'link': serialize_link(link)
    }), 201

@links_bp.route('/<int:id>/attachments/<int:attachment_id>/download', methods=['GET'])
@login_required
@require_permission('links.view')
def download_attachment(id, attachment_id):
    link = db.session.get(Link, id)
    if not link:
        return jsonify({'error': 'Link not found'}), 404
        
    att = db.session.get(LinkAttachment, attachment_id)
    if not att or att.link_id != id:
        return jsonify({'error': 'Attachment not found'}), 404
        
    upload_dir = current_app.config.get('UPLOAD_FOLDER', os.path.join(current_app.config.get('APP_ROOT', '.'), 'data', 'uploads', 'attachments'))
    file_path = os.path.join(upload_dir, att.filename)
    if not os.path.exists(file_path):
        return jsonify({'error': 'Attachment file missing from storage'}), 404
        
    is_view = request.args.get('view') in ('1', 'true', 'yes')
    as_attachment = not is_view
    
    mimetype, _ = mimetypes.guess_type(att.original_filename)
    if not mimetype:
        if att.original_filename.lower().endswith('.pdf'):
            mimetype = 'application/pdf'
        elif att.original_filename.lower().endswith(('.vsd', '.vsdx')):
            mimetype = 'application/vnd.visio'
        else:
            mimetype = 'application/octet-stream'

    return send_from_directory(
        upload_dir,
        att.filename,
        as_attachment=as_attachment,
        download_name=att.original_filename,
        mimetype=mimetype
    )

@links_bp.route('/<int:id>/attachments/<int:attachment_id>', methods=['DELETE'])
@login_required
@require_permission('links.edit')
def delete_attachment(id, attachment_id):
    link = db.session.get(Link, id)
    if not link:
        return jsonify({'error': 'Link not found'}), 404
        
    att = db.session.get(LinkAttachment, attachment_id)
    if not att or att.link_id != id:
        return jsonify({'error': 'Attachment not found'}), 404
        
    orig_name = att.original_filename
    stored_name = att.filename
    db.session.delete(att)
    db.session.commit()
    
    remove_attachment_file_from_disk(stored_name)
    
    write_log('links', 'attachment_deleted', session.get('username', 'system'), link.link_id,
              {'filename': orig_name}, ip_address=request.remote_addr)
              
    return jsonify({'message': 'Attachment deleted', 'link': serialize_link(link)}), 200

@links_bp.route('/<int:id>', methods=['DELETE'])
@login_required
@require_permission('links.delete')
def delete_link(id):
    link = db.session.get(Link, id)
    if not link:
        return jsonify({'error': 'Link not found'}), 404
    link_id = link.link_id
    leg_name = link.leg_name
    mw_ip = link.mw_ip
    
    # Clean up attachment files on disk
    if hasattr(link, 'attachments'):
        for att in link.attachments.all():
            remove_attachment_file_from_disk(att.filename)

    db.session.delete(link)
    db.session.commit()

    # Audit log: link deleted (Tier 2.4)
    write_log('links', 'link_deleted', session.get('username', 'system'), link_id,
              {'leg_name': leg_name, 'mw_ip': mw_ip}, ip_address=request.remote_addr)

    return jsonify({"message": "Link deleted", "link_id": link_id}), 200


@links_bp.route('/<int:id>/ping', methods=['POST'])
@login_required
@require_permission('links.ping')
def manual_ping(id):
    link = db.session.get(Link, id)
    if not link:
        return jsonify({'error': 'Link not found'}), 404
    try:
        result = ping_single_link(link)
        return jsonify({
            "reachable": result.reachable,
            "latency_ms": result.latency_ms,
            "packet_loss": result.packet_loss,
            "timestamp": result.timestamp.isoformat() + "Z"
        }), 200
    except Exception as e:
        return jsonify({"error": str(e)}), 503


@links_bp.route('/legs', methods=['GET'])
@login_required
@require_permission('links.view')
def list_legs():
    legs = db.session.query(Link.leg_name).distinct().order_by(Link.leg_name).all()
    return jsonify({'legs': [row[0] for row in legs]}), 200


@links_bp.route('/<int:id>/metrics', methods=['POST'])
@login_required
@require_permission('links.edit')
def submit_metric(id):
    link = db.session.get(Link, id)
    if not link:
        return jsonify({'error': 'Link not found'}), 404

    data = request.get_json() or {}

    status = LinkStatus.query.filter_by(link_id=link.id).first()
    if not status:
        status = LinkStatus(link_id=link.id)
        db.session.add(status)
        
    status.leg_util_pct = data.get('leg_util_pct', data.get('fiber_util_pct', status.leg_util_pct))
    status.leg_capacity_mbps = data.get('leg_capacity_mbps', data.get('fiber_capacity_mbps', status.leg_capacity_mbps))
    status.mw_util_pct = data.get('mw_util_pct', status.mw_util_pct)
    status.mw_capacity_mbps = data.get('mw_capacity_mbps', status.mw_capacity_mbps)
    status.metric_source = data.get('source', 'manual')
    status.last_metric_at = datetime.utcnow()

    settings = db.session.get(AppSettings, 1)
    warn_pct = settings.util_warning_threshold_pct if settings else 70.0
    crit_pct = settings.util_critical_threshold_pct if settings else 90.0

    mw_util = data.get('mw_util_pct')
    if mw_util is not None:
        if mw_util >= crit_pct:
            send_event_notification('mw_util_high', f'Link {link.link_id} MW utilization at {mw_util:.1f}% (critical)', link_id=link.link_id, severity='critical')
        elif mw_util >= warn_pct:
            send_event_notification('mw_util_high', f'Link {link.link_id} MW utilization at {mw_util:.1f}% (warning)', link_id=link.link_id, severity='warning')

    leg_util = data.get('leg_util_pct', data.get('fiber_util_pct'))
    if leg_util is not None:
        if leg_util >= crit_pct:
            send_event_notification('leg_util_near_cap', f'Link {link.link_id} leg utilization at {leg_util:.1f}% (near capacity)', link_id=link.link_id, severity='critical')
        elif leg_util >= warn_pct:
            send_event_notification('leg_util_high', f'Link {link.link_id} leg utilization at {leg_util:.1f}% (warning)', link_id=link.link_id, severity='warning')

    db.session.commit()
    return jsonify({'result': 'metric recorded'}), 201


@links_bp.route('/bulk-template', methods=['GET'])
@login_required
@require_permission('links.bulk_add')
def download_bulk_template():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bulk Links Template"
    
    headers = [
        "link_id", "leg_name", "site_a", "site_b", 
        "mw_ip", "util_warning_threshold_pct", 
        "util_critical_threshold_pct", "notes"
    ]
    ws.append(headers)
    
    # Add a sample row
    ws.append([
        "LINK-1234", "NorthRegion-Leg1", "SiteA-Router", "SiteB-Router", 
        "192.168.10.5", 75, 95, "Sample row, please delete"
    ])
    
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    return send_file(
        output,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name='bulk_links_template.xlsx'
    )

@links_bp.route('/bulk', methods=['POST'])
@login_required
@require_permission('links.bulk_add')
def bulk_upload():
    if 'file' not in request.files:
        return jsonify({"error": "No file part in the request"}), 400
        
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400
        
    if not file.filename.endswith('.xlsx'):
        return jsonify({"error": "Invalid file type. Please upload an .xlsx file"}), 400

    try:
        wb = openpyxl.load_workbook(file)
        ws = wb.active
    except Exception as e:
        return jsonify({"error": f"Failed to read Excel file: {str(e)}"}), 400

    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        return jsonify({"error": "File is empty or missing headers"}), 400

    headers = [str(h).lower().strip() if h else '' for h in rows[0]]
    
    required_cols = ['link_id', 'leg_name', 'mw_ip']
    for req in required_cols:
        if req not in headers:
            return jsonify({"error": f"Missing required column: {req}"}), 400

    idx_map = {h: i for i, h in enumerate(headers)}
    
    success_count = 0
    failed_rows = []
    
    for row_idx, row in enumerate(rows[1:], start=2):
        try:
            link_id = str(row[idx_map['link_id']]).strip() if row[idx_map.get('link_id')] else None
            leg_name = str(row[idx_map['leg_name']]).strip() if row[idx_map.get('leg_name')] else None
            mw_ip = str(row[idx_map['mw_ip']]).strip() if row[idx_map.get('mw_ip')] else None
            
            if not link_id or not leg_name or not mw_ip:
                failed_rows.append({"row": row_idx, "reason": "Missing link_id, leg_name, or mw_ip"})
                continue
                
            if not validate_ipv4(mw_ip):
                failed_rows.append({"row": row_idx, "reason": "Invalid IPv4 address format"})
                continue

            existing = Link.query.filter_by(link_id=link_id).first()
            if existing:
                failed_rows.append({"row": row_idx, "reason": f"Link ID {link_id} already exists"})
                continue

            site_a = str(row[idx_map['site_a']]).strip() if 'site_a' in idx_map and row[idx_map['site_a']] is not None else None
            site_b = str(row[idx_map['site_b']]).strip() if 'site_b' in idx_map and row[idx_map['site_b']] is not None else None
            notes = str(row[idx_map['notes']]).strip() if 'notes' in idx_map and row[idx_map['notes']] is not None else None
            
            warn_pct = None
            if 'util_warning_threshold_pct' in idx_map and row[idx_map['util_warning_threshold_pct']] is not None:
                try:
                    warn_pct = float(row[idx_map['util_warning_threshold_pct']])
                except ValueError:
                    pass

            crit_pct = None
            if 'util_critical_threshold_pct' in idx_map and row[idx_map['util_critical_threshold_pct']] is not None:
                try:
                    crit_pct = float(row[idx_map['util_critical_threshold_pct']])
                except ValueError:
                    pass

            link = Link(
                link_id=link_id,
                leg_name=leg_name,
                site_a=site_a,
                site_b=site_b,
                mw_ip=mw_ip,
                link_type='microwave',
                notes=notes,
                is_active=True,
                util_warning_threshold_pct=warn_pct,
                util_critical_threshold_pct=crit_pct,
            )
            db.session.add(link)
            db.session.flush() # Ensure link has an ID before fetching external status
            try:
                refresh_external_utilization_for_single_link(link)
            except Exception as e:
                logger.warning(f"Failed to fetch initial external status for bulk link {link_id}: {e}")
                
            success_count += 1
            
        except Exception as e:
            failed_rows.append({"row": row_idx, "reason": str(e)})

    try:
        db.session.commit()
    except Exception as e:
        db.session.rollback()
        return jsonify({"error": f"Database error: {str(e)}"}), 500

    if success_count > 0:
        write_log('links', 'bulk_add_links', session.get('username', 'system'), 'BULK',
                  {'count': success_count, 'failures': len(failed_rows)}, ip_address=request.remote_addr)

    return jsonify({
        "message": f"Successfully imported {success_count} links.",
        "success_count": success_count,
        "failed_count": len(failed_rows),
        "failures": failed_rows
    }), 201
